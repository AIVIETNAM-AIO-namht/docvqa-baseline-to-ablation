"""Cập nhật ô "Cấu hình D" trong Outline §IV từ "chưa chạy" -> kết quả probe.

Chạy 1 lần. Backup tự động sang `.BACKUP-08-10.docx` trước khi sửa.

    python update_outline_config_d.py
"""
import shutil
import sys
from pathlib import Path

import docx
from docx.shared import Pt

BASE = Path(__file__).parent
SRC = BASE / "Outline Project - Hỏi đáp trên ảnh tài liệu (Document VQA).docx"
BAK = BASE / "Outline Project - Hỏi đáp trên ảnh tài liệu (Document VQA).BACKUP-08-10.docx"
sys.stdout.reconfigure(encoding="utf-8")

NEW = (
    "CHƯA DỰNG — probe khả thi zero-shot cho kết luận âm.\n"
    "Qwen2.5-VL-3B zero-shot, 120 câu visual_bold_lookup (dev, Colab T4):\n"
    "ANLS 0,555 · EvF1 0,704 · điểm 0,578 vs Cấu hình B 0,818 (chênh −0,241).\n"
    "Chọn đúng hàng 60/120 = đúng mức đoán bừa; không trả lời được A/B: 0/120.\n"
    "Thang sanity 4 bậc: model chép đúng chữ hai dải nhưng chọn theo VỊ TRÍ\n"
    "(đảo dải → chữ không lật; đổi nhãn → chữ không lật) ⇒ lỗi ở tầng so sánh.\n"
    "⇒ Zero-shot vô dụng. QLoRA là đường duy nhất còn lại, chờ TA xác nhận."
)

OLD = "chưa chạy"


def fill(cell, text):
    cell.text = ""
    for i, line in enumerate(text.split("\n")):
        par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        par.add_run(line).font.size = Pt(8)


def main():
    if not BAK.exists():
        shutil.copy2(SRC, BAK)
        print(f"backup -> {BAK.name}")

    d = docx.Document(SRC)
    from docx.table import Table

    tbl = None
    for i, el in enumerate(d.element.body.iterchildren()):
        if i == 51:
            tbl = Table(el, d)
    assert tbl is not None and len(tbl.rows) == 5, "bảng ablation đổi cấu trúc"

    hits = [c for row in tbl.rows for c in row.cells if c.text.strip() == OLD]
    assert len(hits) == 1, f"thấy {len(hits)} ô '{OLD}', cần đúng 1"
    fill(hits[0], NEW)

    d.save(SRC)
    print(f"OK  cập nhật ô Cấu hình D ({len(NEW)} ký tự)")


if __name__ == "__main__":
    main()
