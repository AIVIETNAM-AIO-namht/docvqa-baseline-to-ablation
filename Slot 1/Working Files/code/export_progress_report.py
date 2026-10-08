"""Export a progress report (markdown) to .docx.

Usage:  python export_progress_report.py ["WEEK02 - Báo cáo tiến độ (22-24.09).md"]
Default source: WEEK01 - Báo cáo tiến độ (16-18.09).md

Handles: h1-h4, tables (with header row), bullets, numbered lists,
blockquotes, inline **bold** and `code`.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(__file__).resolve().parents[1]
NAME = sys.argv[1] if len(sys.argv) > 1 else "WEEK01 - Báo cáo tiến độ (16-18.09).md"
# ponytail: báo cáo nằm trong thư mục con "Báo cáo"; vẫn nhận cả file ở gốc.
SRC = next((p for p in (BASE / "Báo cáo" / NAME, BASE / NAME) if p.exists()),
           BASE / "Báo cáo" / NAME)
DST = SRC.with_suffix(".docx")

# ---------------------------------------------------------------- inline runs

TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+`)")


def add_runs(par, text):
    """Emit **bold**, *italic* and `code` spans as separate runs."""
    for piece in TOKEN.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*"):
            par.add_run(piece[1:-1]).italic = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = par.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xA3, 0x1D, 0x1D)
        else:
            par.add_run(piece)


def shade(cell, hex_fill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(el)


# ---------------------------------------------------------------- md -> docx

def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_sep(line):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def convert(md_path, docx_path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # ---- table
        if stripped.startswith("|") and i + 1 < len(lines) and is_sep(lines[i + 1]):
            header = split_row(stripped)
            body = []
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                body.append(split_row(lines[j]))
                j += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for k, text in enumerate(header):
                cell = table.rows[0].cells[k]
                cell.text = ""
                add_runs(cell.paragraphs[0], text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                shade(cell, "D9E2F3")
            for row in body:
                cells = table.add_row().cells
                for k in range(len(header)):
                    cells[k].text = ""
                    add_runs(cells[k].paragraphs[0], row[k] if k < len(row) else "")
            doc.add_paragraph()
            i = j
            continue

        # ---- heading
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level, text = len(m.group(1)), m.group(2)
            par = doc.add_heading("", level=min(level, 4))
            add_runs(par, text)
            i += 1
            continue

        # ---- horizontal rule
        if re.fullmatch(r"-{3,}", stripped):
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = par.add_run("—" * 20)
            run.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
            i += 1
            continue

        # ---- blockquote
        if stripped.startswith(">"):
            chunk = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                chunk.append(lines[i].strip().lstrip(">").strip())
                i += 1
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Pt(24)
            add_runs(par, " ".join(chunk))
            for run in par.runs:
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            continue

        # ---- bullet
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            par = doc.add_paragraph(style="List Bullet")
            add_runs(par, m.group(1))
            i += 1
            continue

        # ---- numbered
        m = re.match(r"^\d+\.\s+(.*)$", stripped)
        if m:
            par = doc.add_paragraph(style="List Number")
            add_runs(par, m.group(1))
            i += 1
            continue

        # ---- blank
        if not stripped:
            i += 1
            continue

        # ---- paragraph
        par = doc.add_paragraph()
        add_runs(par, stripped)
        i += 1

    doc.save(docx_path)
    return docx_path


if __name__ == "__main__":
    assert SRC.exists(), f"missing source: {SRC}"
    out = convert(SRC, DST)
    size = out.stat().st_size
    print(f"OK  {out.name}")
    print(f"    {size:,} bytes  ({size / 1024:.1f} KB)  -- limit is 10 MB")
    assert size < 10 * 1024 * 1024, "over the 10 MB form limit"
