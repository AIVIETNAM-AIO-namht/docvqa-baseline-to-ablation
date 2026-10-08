"""Kiểm tra nhanh file .docx vừa export: có đủ mục, bảng không lệch cột."""
import sys
from pathlib import Path

import docx

sys.stdout.reconfigure(encoding="utf-8")
BASE = Path(__file__).resolve().parents[1] / "Báo cáo"

for name in ("WEEK03 - Báo cáo tiến độ.docx", "WEEK03 - Keeptrack.docx"):
    d = docx.Document(BASE / name)
    h1 = [p.text for p in d.paragraphs if p.style.name == "Heading 1"]
    print(f"{name}\n  H1: {h1}")
    for t in d.tables:
        widths = {len(r.cells) for r in t.rows}
        assert len(widths) == 1, f"  !! bảng lệch cột: {widths}"
    print(f"  bảng: {len(d.tables)} — không bảng nào lệch cột")
